import io
import openpyxl
from datetime import datetime
import streamlit as st
from weasyprint import HTML

def get_val(ws, r, c):
    return ws.cell(row=r, column=c).value

def format_date(val):
    if isinstance(val, datetime):
        return val.strftime('%d/%m/%Y')
    return str(val) if val else "Não informado"

def format_money(val):
    if val is None or val == "":
        return "R$ 0,00"
    try:
        v = float(val)
        prefix = "-" if v < 0 else ""
        return f"{prefix}R$ {abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except:
        return "R$ 0,00"

st.set_page_config(page_title="Termo de Encerramento", page_icon="📄", layout="centered")

st.title("📄 Gerador de Termo de Encerramento de Contrato")
st.write("Faça o upload da planilha de encerramento para calcular débitos/reembolsos e gerar o PDF de acerto.")

file_excel = st.file_uploader("Selecione a Planilha de Encerramento (.xlsx)", type=["xlsx"])

if file_excel:
    if st.button("🚀 Gerar Termo em PDF", type="primary"):
        with st.spinner("Processando cálculos e gerando o documento..."):
            wb = openpyxl.load_workbook(file_excel, data_only=True)
            ws = wb["Planilha1"]
            
            # Dados principais
            locador = get_val(ws, 4, 2) or "Não informado"
            locatario = get_val(ws, 5, 2) or "Não informado"
            imovel = get_val(ws, 6, 2) or "Não informado"
            iptu_num = get_val(ws, 7, 2) or "Não informado"
            
            aluguel = float(get_val(ws, 9, 2) or 0.0)
            dt_inicio_raw = get_val(ws, 14, 2)
            dt_fim_raw = get_val(ws, 15, 2)
            dt_rescisao_raw = get_val(ws, 15, 5)
            
            # 1. Multa Rescisória
            tem_multa_raw = get_val(ws, 17, 2)
            flag_multa = str(tem_multa_raw).strip().upper() in ['1', 'S', 'SIM', 'TRUE']
            multa_calc = 0.0
            if flag_multa and isinstance(dt_inicio_raw, datetime) and isinstance(dt_fim_raw, datetime) and isinstance(dt_rescisao_raw, datetime):
                prazo_total = (dt_fim_raw - dt_inicio_raw).days
                tempo_decorrido = (dt_rescisao_raw - dt_inicio_raw).days
                dias_restantes = prazo_total - tempo_decorrido
                if prazo_total > 0 and dias_restantes > 0:
                    multa_calc = round(((aluguel * 3) / prazo_total) * dias_restantes, 2)

            # 2. Seguro Incêndio (Reembolso)
            dt_incendio_inicio = get_val(ws, 12, 3)
            vlr_seguro_total = float(get_val(ws, 12, 5) or 0.0)
            reembolso_seguro_flag = str(get_val(ws, 13, 3)).strip().upper() in ['S', 'SIM', 'TRUE', '1']
            
            seguro_reembolso_calc = 0.0
            if reembolso_seguro_flag and isinstance(dt_incendio_inicio, datetime) and isinstance(dt_rescisao_raw, datetime) and vlr_seguro_total > 0:
                dias_efetivos = (dt_rescisao_raw - dt_incendio_inicio).days + 1
                if dias_efetivos > 0:
                    seguro_reembolso_calc = round((vlr_seguro_total / 365.0) * dias_efetivos * 0.8025, 2)

            # Extração dos Lançamentos Financeiros da Planilha
            itens_financeiros = []
            total_debitos = 0.0
            
            for r in range(32, 42):
                nome = get_val(ws, r, 1)
                val = get_val(ws, r, 3)
                
                # Ajuste da Multa na linha 36 se recalculada
                if r == 36 and flag_multa:
                    val = multa_calc
                
                # Ajuste do Seguro Incêndio na linha 37 (Crédito para o locatário se houver reembolso)
                if r == 37 and reembolso_seguro_flag and seguro_reembolso_calc > 0:
                    val = -seguro_reembolso_calc
                
                if val is not None and val != "" and float(val or 0) != 0 and nome is not None:
                    v_float = float(val)
                    total_debitos += v_float
                    itens_financeiros.append({
                        "nome": nome,
                        "valor_float": v_float,
                        "valor_str": format_money(v_float)
                    })

            # Abatimento de Caução
            valor_caucao = float(get_val(ws, 42, 3) or 0.0) if ws.max_row >= 42 else 0.0
            saldo_final = total_debitos - valor_caucao
            
            if saldo_final > 0:
                texto_total = "VALOR A SER COBRADO DO LOCATÁRIO:"
                valor_total_str = format_money(saldo_final)
                cor_total = "#c0392b"
            elif saldo_final < 0:
                texto_total = "VALOR A DEVOLVER AO LOCATÁRIO:"
                valor_total_str = format_money(abs(saldo_final))
                cor_total = "#27ae60"
            else:
                texto_total = "ACERTO QUITADO:"
                valor_total_str = "R$ 0,00"
                cor_total = "#2c3e50"

            # Estrutura HTML do PDF
            html_content = f"""
            <!DOCTYPE html>
            <html>
            <head>
            <meta charset="utf-8">
            <style>
                @page {{ size: A4; margin: 20mm 15mm; background-color: #ffffff; }}
                body {{ font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif; color: #333; margin: 0; padding: 0; font-size: 11pt; }}
                .header {{ text-align: center; border-bottom: 2px solid #2c3e50; padding-bottom: 10px; margin-bottom: 20px; }}
                .header h1 {{ margin: 0; color: #2c3e50; font-size: 18pt; text-transform: uppercase; }}
                .header p {{ margin: 5px 0 0 0; font-size: 10pt; color: #7f8c8d; }}
                .section {{ margin-bottom: 20px; }}
                .section-title {{ font-size: 11pt; font-weight: bold; color: #fff; background-color: #2c3e50; padding: 5px 10px; margin-bottom: 10px; text-transform: uppercase; }}
                table {{ width: 100%; border-collapse: collapse; margin-bottom: 10px; }}
                th, td {{ padding: 8px; border: 1px solid #bdc3c7; text-align: left; }}
                th {{ background-color: #ecf0f1; font-weight: bold; width: 30%; }}
                .table-valores th, .table-valores td {{ border: none; border-bottom: 1px solid #ecf0f1; }}
                .table-valores td.valor {{ text-align: right; font-weight: bold; }}
                .total-box {{ margin-top: 15px; padding: 12px; background-color: #ecf0f1; border-left: 6px solid {cor_total}; font-size: 13pt; font-weight: bold; text-align: right; }}
                .signatures {{ margin-top: 50px; width: 100%; text-align: center; }}
                .sig-line {{ display: inline-block; width: 45%; border-top: 1px solid #333; margin: 0 2%; padding-top: 5px; }}
            </style>
            </head>
            <body>
                <div class="header">
                    <h1>Termo de Encerramento de Contrato</h1>
                    <p>Resumo de Acerto Financeiro e Devolução de Imóvel</p>
                </div>

                <div class="section">
                    <div class="section-title">Dados do Contrato</div>
                    <table>
                        <tr><th>Locador:</th><td>{locador}</td></tr>
                        <tr><th>Locatário:</th><td>{locatario}</td></tr>
                        <tr><th>Imóvel:</th><td>{imovel}</td></tr>
                        <tr><th>Inscrição IPTU/TLP:</th><td>{iptu_num}</td></tr>
                    </table>
                </div>

                <div class="section">
                    <div class="section-title">Prazos e Datas</div>
                    <table>
                        <tr><th>Início do Contrato:</th><td>{format_date(dt_inicio_raw)}</td></tr>
                        <tr><th>Fim do Contrato:</th><td>{format_date(dt_fim_raw)}</td></tr>
                        <tr><th>Data Rescisão/Chaves:</th><td>{format_date(dt_rescisao_raw)}</td></tr>
                    </table>
                </div>

                <div class="section">
                    <div class="section-title">Apuração Financeira para Acerto</div>
                    <table class="table-valores">
            """
            
            for item in itens_financeiros:
                html_content += f"<tr><td>{item['nome']}</td><td class='valor'>{item['valor_str']}</td></tr>"

            if valor_caucao > 0:
                html_content += f"<tr><td>(-) Abatimento de Caução Depositada</td><td class='valor'>-{format_money(valor_caucao)}</td></tr>"

            html_content += f"""
                    </table>
                    <div class="total-box">{texto_total} {valor_total_str}</div>
                </div>

                <div style="font-size: 8.5pt; color: #7f8c8d; margin-top: 25px; text-align: justify;">
                    Declaro para os devidos fins que as chaves do imóvel acima citado foram entregues nesta data, 
                    estando as partes de acordo com os valores apurados neste termo para a quitação final das 
                    obrigações financeiras inerentes ao contrato de locação.
                </div>

                <div class="signatures">
                    <div class="sig-line"><b>{locatario}</b><br>Locatário</div>
                    <div class="sig-line"><b>{locador}</b><br>Locador (ou Representante)</div>
                </div>
            </body>
            </html>
            """
            
            pdf_buffer = io.BytesIO()
            HTML(string=html_content).write_pdf(pdf_buffer)
            pdf_buffer.seek(0)
            
            nome_arquivo = f"Termo_Encerramento_{str(locatario).replace(' ', '_')}.pdf"

        st.success("✅ Termo em PDF gerado com sucesso!")
        
        st.download_button(
            label="📥 Baixar Termo de Encerramento (PDF)",
            data=pdf_buffer,
            file_name=nome_arquivo,
            mime="application/pdf"
        )
